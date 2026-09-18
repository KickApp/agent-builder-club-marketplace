#!/usr/bin/env python3
"""
Weekly cash forecast builder. Config-driven and date-driven.

Fill CONFIG from the books, run it, verify it. Everything below CONFIG derives the
schedule from dates, so the model re-anchors itself correctly whenever it is rerun.

    python3 build_cash_model.py                     # writes CONFIG["out"]

What makes this a model rather than a table:
  * every grid cell is a formula, never a value, and the schedule lives in the sheet
  * amount, how often, and the start and end dates are all editable cells, so a line can
    be re-timed, re-sized, or switched off without rebuilding anything
  * the week columns are dates derived from one cell, so moving the forecast start moves
    the whole calendar and every schedule re-anchors to it
  * the forecast shows what happens with no rescue, so the dip is the real dip and the
    shortfall tile is the whole funding answer
  * two tabs, because a tool that needs a manual is a tool that is too complicated

Read reference/MODELS.md before changing the structure.
"""
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# ==========================================================================================
# CONFIG. Fill from the books. Cadences are computed from the dates you give here.
#
#   cadence: "weekly"   every week
#            "biweekly" every 14 days from "anchor" (any real pay date)
#            "monthly"  the week containing "day" of each month (default 1)
#            "dates"    an explicit list in "dates"
# ==========================================================================================
CONFIG = {
    "entity": "Northwind Labs Inc",
    "prepared_note": "built from the cash ledger",
    "out": "cash-forecast.xlsx",

    "run_date": "2026-08-10",        # Monday that anchors week 1
    "weeks": 13,                     # 13 is the standard; 6 to 8 for a tight-cash situation
    "through_year_end": False,

    "start_cash": 62_000,
    "floor": 30_000,

    "inflows": [
        {"label": "Collections on receivables", "amount": 36_000, "cadence": "weekly"},
        {"label": "New sales receipts", "amount": 4_000, "cadence": "weekly"},
    ],
    "outflows": [
        {"label": "Payroll and payroll taxes", "amount": 38_400,
         "cadence": "biweekly", "anchor": "2026-08-14"},
        {"label": "Vendors and operating spend", "amount": 17_000, "cadence": "weekly"},
        {"label": "Rent", "amount": 12_000, "cadence": "monthly", "day": 1},
        {"label": "Loan payments", "amount": 6_400, "cadence": "monthly", "day": 1},
        {"label": "Owner draw", "amount": 12_000, "cadence": "monthly", "day": 10},
        {"label": "Insurance renewal", "amount": 9_000,
         "cadence": "dates", "dates": ["2026-09-15"]},
    ],

    # Assumptions tab: label -> where the number came from
    "bases": {
        "Collections on receivables": "AR aging, 61 days to collect",
        "New sales receipts": "Trailing 13-week average",
        "Payroll and payroll taxes": "Payroll register, current headcount",
        "Vendors and operating spend": "Recurring vendor list, trailing 13 weeks",
        "Rent": "Lease",
        "Loan payments": "Amortization schedule",
        "Owner draw": "User-stated, monthly",
        "Insurance renewal": "Renewal notice",
    },
}

# ==========================================================================================
# DERIVE the schedule from dates
# ==========================================================================================
CAD_WEEKLY, CAD_BIWEEK, CAD_MONTH, CAD_ONCE = (
    "Every week", "Every 2 weeks", "Monthly", "One time")
CADENCES = [CAD_WEEKLY, CAD_BIWEEK, CAD_MONTH, CAD_ONCE]


def D(s):
    return date.fromisoformat(s)


run_date = D(CONFIG["run_date"])
if CONFIG.get("through_year_end"):
    WK = ((date(run_date.year, 12, 31) - run_date).days // 7) + 1
else:
    WK = int(CONFIG["weeks"])
weeks = [run_date + timedelta(days=7 * i) for i in range(WK)]
horizon_end = weeks[-1] + timedelta(days=6)


def week_of(dt):
    """Which grid column does this date land in? None if outside the horizon."""
    for i, w in enumerate(weeks):
        if w <= dt <= w + timedelta(days=6):
            return i
    return None


def months_in_horizon():
    out, yy, mm = [], run_date.year, run_date.month
    while date(yy, mm, 1) <= horizon_end:
        out.append((yy, mm))
        mm += 1
        if mm > 12:
            mm, yy = 1, yy + 1
    return out


def schedule(spec):
    """The set of column indexes this line hits. 'weekly' returns the string 'weekly'.

    The builder does not use this. It exists so model_test.py can work out the answer a
    second way, in Python, and hold the sheet's formulas against it. Two implementations
    that agree are worth more than one that is merely self-consistent.
    """
    cad = spec.get("cadence", "weekly")
    if cad == "weekly":
        return "weekly"
    hits = set()
    if cad == "biweekly":
        p = D(spec["anchor"])
        while p > weeks[0]:
            p -= timedelta(days=14)
        while p <= horizon_end:
            i = week_of(p)
            if i is not None:
                hits.add(i)
            p += timedelta(days=14)
    elif cad == "monthly":
        day = int(spec.get("day", 1))
        for (yy, mm) in months_in_horizon():
            i = week_of(date(yy, mm, min(day, 28)))
            if i is not None:
                hits.add(i)
    elif cad == "dates":
        for s in spec.get("dates", []):
            i = week_of(D(s))
            if i is not None:
                hits.add(i)
    else:
        raise ValueError(f"unknown cadence {cad!r} on {spec['label']!r}")
    return hits


def expand(lines):
    """One row per scheduled thing.

    A spec listing three dates used to be one row that lit up three columns, which read as
    three hardcoded numbers because that is what it was. Each date becomes its own row with
    its own editable date cell, so "what if that one slips two weeks" is a cell edit.
    """
    out = []
    for spec in lines:
        dates = spec.get("dates") or []
        if spec.get("cadence") == "dates" and len(dates) > 1:
            for s in dates:
                one = dict(spec, dates=[s],
                           label=f"{spec['label']} ({D(s).strftime('%b %-d')})")
                out.append(one)
        else:
            out.append(dict(spec))
    return out


def sheet_schedule(spec):
    """(cadence label, first date, last date) as the three cells the sheet will carry."""
    cad = spec.get("cadence", "weekly")
    last = D(spec["until"]) if spec.get("until") else None
    if cad == "weekly":
        return CAD_WEEKLY, (D(spec["from"]) if spec.get("from") else None), last
    if cad == "biweekly":
        return CAD_BIWEEK, D(spec["anchor"]), last
    if cad == "monthly":
        # The day of the month is carried by a real date, so the reader edits a date rather
        # than an integer whose meaning they have to infer.
        day = min(int(spec.get("day", 1)), 28)
        return CAD_MONTH, date(run_date.year, run_date.month, day), last
    if cad == "dates":
        return CAD_ONCE, D(spec["dates"][0]), last
    raise ValueError(f"unknown cadence {cad!r} on {spec['label']!r}")


def cadence_text(spec):
    """Plain English for the HTML companion, which has no cells to point at.

    Derived from sheet_schedule rather than from the raw config, so the sentence and the
    workbook cells can never describe different things.
    """
    cad, frm, until = sheet_schedule(spec)
    d = lambda x: x.strftime("%b %-d")
    if cad == CAD_ONCE:
        base = f"Once, on {d(frm)}"
    elif cad == CAD_BIWEEK:
        base = f"Every 2 weeks from {d(frm)}"
    elif cad == CAD_MONTH:
        day = frm.day
        suffix = "th" if 11 <= day <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(day % 10, "th")
        base = f"Monthly, around the {day}{suffix}"
    else:
        base = f"Every week from {d(frm)}" if frm else "Every week"
    return base + (f", until {d(until)}" if until else "")


# ==========================================================================================
# STYLE. Charcoal and off-white, matching the HTML artifacts rather than shouting
# ==========================================================================================
FONT = "Arial"
MONEY = '$#,##0;($#,##0);"-"'
DATEF = "mmm d"
C_AMT, C_CAD, C_FROM, C_UNTIL = 2, 3, 4, 5    # the editable schedule
C0 = 6                                        # first week column
TOT = C0 + WK                                 # totals column index
TOTL = get_column_letter(TOT)
LASTW = get_column_letter(C0 + WK - 1)

INK = "252533"
MUT = "6B6B7B"


def f(size=9, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


f_ent = f(9, False, MUT)
f_title = f(15, True, INK)
f_sub = f(9, False, MUT)
f_sect = f(10, True, INK)
f_hdr = f(9, True, INK)
f_norm, f_bold = f(9), f(9, True)
f_inp = f(9, False, "0B4FBF")                 # blue = a number you can type over
f_note = f(8, False, MUT, True)
f_tilelab = f(8, False, MUT)
f_tileval = f(14, True, INK)

soft = PatternFill("solid", fgColor="F4F4F7")     # section band
edit = PatternFill("solid", fgColor="FFF3CD")     # editable cell
tile = PatternFill("solid", fgColor="F7F7FA")
negfill = PatternFill("solid", fgColor="F9E4E4")  # pastel red, never alarm red
thin = Side(style="thin", color="E3E3EA")
rule = Side(style="thin", color="C9C9D2")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
topline = Border(top=rule)


def col(i):
    return get_column_letter(C0 + i)


R_HEAD = 10                                   # the week-date header row


def hit_formula(rr, i):
    """Does the line on row `rr` pay out in the week starting at this column's date?

    The whole point of this function is that the answer is computed by the spreadsheet from
    cells the reader can edit, not by Python at build time. Every branch reads $C, $D or $E
    on the row, and the week's own date out of the header, so changing any of the four moves
    the money.
    """
    W = f"{col(i)}${R_HEAD}"
    amt, cad, frm, until = f"$B{rr}", f"$C{rr}", f"$D{rr}", f"$E{rr}"
    # A monthly line lands in this week if its day of the month falls inside it. Two
    # candidates, because a week straddles a month boundary about a quarter of the time.
    mday = f"MIN(DAY({frm}),28)"
    occ1 = f"DATE(YEAR({W}),MONTH({W}),{mday})"
    occ2 = f"DATE(YEAR({W}+6),MONTH({W}+6),{mday})"
    monthly = f"OR(AND({occ1}>={W},{occ1}<={W}+6),AND({occ2}>={W},{occ2}<={W}+6))"
    # MOD is measured from the anchor forward, so an anchor later in the week still counts.
    biweekly = f"MOD({frm}-{W},14)<=6"
    once = f"AND({frm}>={W},{frm}<={W}+6)"
    # Blank start or end means "no bound", which is what a reader expects an empty cell to do.
    window = f'OR({frm}="",{W}+6>={frm}),OR({until}="",{W}<={until})'
    cadence = (f'IF({cad}="{CAD_WEEKLY}",TRUE,'
               f'IF({cad}="{CAD_BIWEEK}",{biweekly},'
               f'IF({cad}="{CAD_MONTH}",{monthly},'
               f'IF({cad}="{CAD_ONCE}",{once},FALSE))))')
    return f"=IF(AND({window},{cadence}),{amt},0)"


def build(path=None):
    path = path or CONFIG["out"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"
    ws.sheet_view.showGridLines = False

    inflows, outflows = expand(CONFIG["inflows"]), expand(CONFIG["outflows"])
    floor_v, start_v = CONFIG["floor"], CONFIG["start_cash"]

    # ---------------------------------------------------------------- title
    ws["A1"], ws["A1"].font = CONFIG["entity"], f_ent
    ws["A2"], ws["A2"].font = f"{WK}-week cash forecast", f_title
    ws["A3"].value = (f"{weeks[0].strftime('%b %-d')} to {horizon_end.strftime('%b %-d, %Y')}"
                      f" · {CONFIG['prepared_note']}")
    ws["A3"].font = f_sub

    # ---------------------------------------------------------------- input cells
    # Declared before the tiles so the tile formulas can point at them.
    r_start, r_floor, r_wk1 = 5, 6, 7   # left of the tiles, so nothing overlaps the subtitle
    for r, label, val, fmt_ in ((r_start, "Starting cash", start_v, MONEY),
                                (r_floor, "Minimum balance", floor_v, MONEY),
                                (r_wk1, "Week 1 starts", weeks[0], DATEF)):
        ws.cell(r, 1, label).font = f_norm
        c = ws.cell(r, 2, val)
        c.font, c.fill, c.border, c.number_format = f_inp, edit, box, fmt_
    B_START, B_FLOOR = f"$B${r_start}", f"$B${r_floor}"

    # A dropdown rather than free text, because the grid formulas match on these exact
    # strings and a typo would silently zero a whole row.
    dv_cad = DataValidation(type="list", formula1='"' + ",".join(CADENCES) + '"',
                            allow_blank=False, showDropDown=False)
    dv_cad.error = "Pick one of: " + ", ".join(CADENCES)
    dv_cad.errorTitle = "Not a cadence the model knows"
    ws.add_data_validation(dv_cad)

    # ---------------------------------------------------------------- grid
    r_head = R_HEAD
    for cc, label in ((1, "Week of"), (C_AMT, "Amount"), (C_CAD, "How often"),
                      (C_FROM, "Starting"), (C_UNTIL, "Ending")):
        ws.cell(r_head, cc, label).font = f_hdr
    # Week 1 reads the input cell and every later week is the one before it plus seven, so
    # moving the start date moves all thirteen columns and every schedule with them.
    for i in range(WK):
        c = ws.cell(r_head, C0 + i,
                    f"=$B${r_wk1}" if i == 0 else f"={col(i - 1)}${r_head}+7")
        c.font, c.number_format, c.alignment = f_hdr, DATEF, Alignment(horizontal="center")
    ws.cell(r_head, TOT, "Total").font = f_hdr
    for cc in range(1, TOT + 1):
        ws.cell(r_head, cc).border = Border(bottom=rule)

    r = r_head + 1
    r_open = r
    ws.cell(r, 1, "Starting cash").font = f_norm
    for i in range(WK):
        # week 1 opens at the input cell; later weeks open where the prior week closed
        ws.cell(r, C0 + i, f"={B_START}" if i == 0 else f"={col(i - 1)}{{END}}")
    r += 1

    def section(title, lines, start_row):
        rr = start_row
        ws.cell(rr, 1, title).font = f_sect
        for cc in range(1, TOT + 1):
            ws.cell(rr, cc).fill = soft
        rr += 1
        rows = []
        for spec in lines:
            ws.cell(rr, 1, spec["label"]).font = f_norm
            cad, frm, until = sheet_schedule(spec)
            for cc, val, numfmt in ((C_AMT, spec["amount"], MONEY),
                                    (C_CAD, cad, "@"),
                                    (C_FROM, frm, DATEF),
                                    (C_UNTIL, until, DATEF)):
                c = ws.cell(rr, cc, val)
                c.font, c.fill, c.border, c.number_format = f_inp, edit, box, numfmt
            dv_cad.add(ws.cell(rr, C_CAD))
            for i in range(WK):
                cell = ws.cell(rr, C0 + i, hit_formula(rr, i))
                cell.number_format = MONEY
            rows.append(rr)
            rr += 1
        return rows, rr

    in_rows, r = section("Money in", inflows, r)
    r_tin = r
    ws.cell(r, 1, "Total money in").font = f_bold
    for i in range(WK):
        c = ws.cell(r, C0 + i, "=" + "+".join(f"{col(i)}{x}" for x in in_rows))
        c.font, c.number_format, c.border = f_bold, MONEY, topline
    ws.cell(r, 1).border = topline
    r += 2

    out_rows, r = section("Money out", outflows, r)
    r_tout = r
    ws.cell(r, 1, "Total money out").font = f_bold
    for i in range(WK):
        c = ws.cell(r, C0 + i, "=" + "+".join(f"{col(i)}{x}" for x in out_rows))
        c.font, c.number_format, c.border = f_bold, MONEY, topline
    ws.cell(r, 1).border = topline
    r += 2

    r_net = r
    ws.cell(r, 1, "Net change").font = f_bold
    for i in range(WK):
        c = ws.cell(r, C0 + i, f"={col(i)}{r_tin}-{col(i)}{r_tout}")
        c.font, c.number_format = f_bold, MONEY
    r += 1

    r_end = r
    ws.cell(r, 1, "Ending cash").font = f_bold
    for i in range(WK):
        c = ws.cell(r, C0 + i, f"={col(i)}{r_open}+{col(i)}{r_net}")
        c.font, c.number_format = f_bold, MONEY
    r += 1

    # the opening row could only be written once the ending row had a number
    for i in range(1, WK):
        ws.cell(r_open, C0 + i).value = f"={col(i - 1)}{r_end}"
    ws.cell(r_open, C0).value = f"={B_START}"
    for i in range(WK):
        ws.cell(r_open, C0 + i).number_format = MONEY

    r_stat = r
    ws.cell(r, 1, "Status").font = f_norm
    for i in range(WK):
        c = ws.cell(r, C0 + i,
                    f'=IF({col(i)}{r_end}<0,"Out of cash",IF({col(i)}{r_end}<{B_FLOOR},"Below minimum",""))')
        c.font, c.alignment = f(8, False, "9A4B4B"), Alignment(horizontal="center")
    r += 1

    # totals column
    for rr in (r_tin, r_tout, r_net):
        c = ws.cell(rr, TOT, f"=SUM({col(0)}{rr}:{LASTW}{rr})")
        c.font, c.number_format, c.border = f_bold, MONEY, topline
    for rr in in_rows + out_rows:
        c = ws.cell(rr, TOT, f"=SUM({col(0)}{rr}:{LASTW}{rr})")
        c.font, c.number_format = f_norm, MONEY

    # a hidden helper beats an array formula: the checker evaluates it and Excel never
    # shows the user a CSE formula they might break
    r_help = r + 1
    ws.cell(r_help, 1, "helper: first week below the minimum").font = f_note
    for i in range(WK):
        ws.cell(r_help, C0 + i, f"=IF({col(i)}{r_end}<{B_FLOOR},{i + 1},9999)")
    ws.row_dimensions[r_help].hidden = True

    END_RANGE = f"{col(0)}{r_end}:{LASTW}{r_end}"
    HELP_RANGE = f"{col(0)}{r_help}:{LASTW}{r_help}"
    DATE_RANGE = f"{col(0)}{r_head}:{LASTW}{r_head}"

    # ---------------------------------------------------------------- tiles
    tiles = [
        ("Cash today", f"={B_START}", MONEY),
        ("Lowest cash", f"=MIN({END_RANGE})", MONEY),
        ("Week it happens",
         f'=TEXT(INDEX({DATE_RANGE},MATCH(MIN({END_RANGE}),{END_RANGE},0)),"mmm d")', "@"),
        ("First shortfall",
         f'=IF(MIN({HELP_RANGE})=9999,"Never",'
         f'TEXT(INDEX({DATE_RANGE},MIN({HELP_RANGE})),"mmm d"))', "@"),
        ("Cash needed", f"=MAX(0,{B_FLOOR}-MIN({END_RANGE}))", MONEY),
    ]
    for n, (label, formula, numfmt) in enumerate(tiles):
        c1 = C0 + n * 2                       # aligned to the grid, so every tile is equal width
        lc = ws.cell(5, c1, label)
        lc.font, lc.alignment = f_tilelab, Alignment(horizontal="center")
        vc = ws.cell(6, c1, formula)
        vc.font, vc.number_format = f_tileval, numfmt
        vc.alignment = Alignment(horizontal="center")
        for rr in (5, 6):
            for cc in (c1, c1 + 1):
                ws.cell(rr, cc).fill = tile
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c1 + 1)
        ws.merge_cells(start_row=6, start_column=c1, end_row=6, end_column=c1 + 1)

    ws.cell(8, 1, "Amber cells are yours to change: the amount, how often, and the dates it "
                  "runs between. Everything else is a formula.").font = f_note

    # ---------------------------------------------------------------- conditional format
    for rng in (END_RANGE,):
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="lessThan", formula=[B_FLOOR], fill=negfill, font=f(9, True, "9A4B4B")))

    # ---------------------------------------------------------------- chart
    ch = LineChart()
    ch.title = "Ending cash by week"
    ch.height, ch.width = 7.5, 26
    ch.y_axis.numFmt = '$#,##0'
    data = Reference(ws, min_col=C0, max_col=C0 + WK - 1, min_row=r_end, max_row=r_end)
    ch.add_data(data, from_rows=True, titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=C0, max_col=C0 + WK - 1,
                                min_row=r_head, max_row=r_head))
    ch.series[0].tx = None
    ws.add_chart(ch, f"A{r_help + 2}")

    # ---------------------------------------------------------------- widths
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions[get_column_letter(C_CAD)].width = 14
    ws.column_dimensions[get_column_letter(C_FROM)].width = 10
    ws.column_dimensions[get_column_letter(C_UNTIL)].width = 10
    for i in range(WK):
        ws.column_dimensions[col(i)].width = 11
    ws.column_dimensions[TOTL].width = 13
    ws.freeze_panes = ws.cell(r_head + 1, C0)

    # ================================================================== Assumptions
    a = wb.create_sheet("Assumptions")
    a.sheet_view.showGridLines = False
    a["A1"], a["A1"].font = CONFIG["entity"], f_ent
    a["A2"], a["A2"].font = "Assumptions", f_title
    a["A3"].value = "Every number in the forecast, where it came from, and the cell that controls it."
    a["A3"].font = f_sub

    ar = 5
    for h, w in zip(["Line", "Amount", "How often", "Where it came from", "Cell to change"],
                    [30, 13, 34, 40, 20]):
        c = a.cell(ar, ["Line", "Amount", "How often", "Where it came from",
                        "Cell to change"].index(h) + 1, h)
        c.font, c.border = f_hdr, Border(bottom=rule)
        a.column_dimensions[get_column_letter(
            ["Line", "Amount", "How often", "Where it came from",
             "Cell to change"].index(h) + 1)].width = w
    ar += 1

    def block(title, lines, rows, start):
        rr = start
        a.cell(rr, 1, title).font = f_sect
        for cc in range(1, 6):
            a.cell(rr, cc).fill = soft
        rr += 1
        for spec, grid_row in zip(lines, rows):
            a.cell(rr, 1, spec["label"]).font = f_norm
            v = a.cell(rr, 2, f"=Forecast!$B${grid_row}")
            v.font, v.number_format = f_norm, MONEY
            # Formulas, not the build-time string: this column has to still be true after
            # someone changes the dropdown on the Forecast tab.
            hw = a.cell(rr, 3, f'=Forecast!$C${grid_row}&IF(Forecast!$D${grid_row}="","",'
                               f'" from "&TEXT(Forecast!$D${grid_row},"mmm d"))'
                               f'&IF(Forecast!$E${grid_row}="","",'
                               f'" until "&TEXT(Forecast!$E${grid_row},"mmm d"))')
            hw.font = f_norm
            a.cell(rr, 4, CONFIG["bases"].get(spec["label"].split(" (")[0], "")).font = f_norm
            a.cell(rr, 5, f"Forecast!B{grid_row}:E{grid_row}").font = f(9, False, MUT)
            rr += 1
        return rr + 1

    a.cell(ar, 1, "Balances").font = f_sect
    for cc in range(1, 6):
        a.cell(ar, cc).fill = soft
    ar += 1
    for label, grid_row, basis in (("Starting cash", r_start, "Cash ledger, confirm against the bank"),
                                   ("Minimum balance", r_floor, "Set by the business")):
        a.cell(ar, 1, label).font = f_norm
        v = a.cell(ar, 2, f"=Forecast!$B${grid_row}")
        v.font, v.number_format = f_norm, MONEY
        a.cell(ar, 3, "One-off").font = f_norm
        a.cell(ar, 4, basis).font = f_norm
        a.cell(ar, 5, f"Forecast!B{grid_row}").font = f(9, False, MUT)
        ar += 1
    ar += 1

    ar = block("Money in", inflows, in_rows, ar)
    ar = block("Money out", outflows, out_rows, ar)  # already expanded above

    a.cell(ar, 1, "Forward revenue is an assumption, not booked activity. Every line can be "
                  "re-sized, re-timed, or ended early from its own row on the Forecast "
                  "tab.").font = f_note

    wb.save(path)
    return {"path": path, "ending_row": r_end, "helper_row": r_help,
            "start_cell": f"Forecast!B{r_start}", "floor_cell": f"Forecast!B{r_floor}",
            "tiles_row": 6, "weeks": WK}


if __name__ == "__main__":
    info = build()
    print(f"wrote {info['path']}  ({info['weeks']} weeks)")
    for k, v in info.items():
        print(f"   {k}: {v}")
