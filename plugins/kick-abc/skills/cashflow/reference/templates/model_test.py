#!/usr/bin/env python3
"""Prove the cash model is a model.

The failure this guards against is a workbook that looks formula-driven and is not: a grid
where Python decided which weeks got paid and wrote the answer in, so the file opens fine,
ties out fine, and does nothing at all when the reader changes an input. That is a printed
table with extra steps, and it is invisible until someone tries to use it.

So every check here edits a cell, recalculates the whole workbook with a real formula
engine, and asserts the numbers moved the way a person would expect.

    pip install formulas openpyxl
    python3 model_test.py
"""
import sys
import warnings
from datetime import date, timedelta

warnings.filterwarnings("ignore")

try:
    import formulas
except ImportError:
    print("model_test: formulas is not installed, skipping. Run `pip install formulas`.")
    sys.exit(0)

from openpyxl import load_workbook
import build_cash_model as M

WORK = "/tmp/model_test.xlsx"
fails = []


def recalc(path):
    """Every value in the workbook, keyed as SHEET!A1."""
    sol = formulas.ExcelModel().loads(path).finish().calculate()
    out = {}
    for k, v in sol.items():
        if "]" not in k or "!" not in k:
            continue
        sheet, _, ref = k.upper().rpartition("'!")
        sheet = sheet.rpartition("]")[2]
        try:
            out[f"{sheet}!{ref}"] = v.value[0, 0]
        except Exception:
            pass
    return out


def num(vals, ref):
    v = vals.get(f"FORECAST!{ref}")
    try:
        return round(float(v))
    except (TypeError, ValueError):
        return v


def row_of(ws, label):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == label:
            return r
    raise AssertionError(f"no row labelled {label!r}")


def week_cells(rr):
    return [f"{M.col(i)}{rr}" for i in range(M.WK)]


def build_and_read(edits=None):
    """Write a fresh workbook, apply cell edits, recalculate, return the values."""
    M.build(WORK)
    if edits:
        wb = load_workbook(WORK)
        ws = wb["Forecast"]
        for ref, val in edits(ws).items():
            ws[ref] = val
        wb.save(WORK)
    return recalc(WORK)


def check(name, ok, detail=""):
    print(f"  {'pass' if ok else 'FAIL'}  {name}")
    if not ok:
        fails.append(name)
        if detail:
            print(f"        {detail}")


# ---------------------------------------------------------------- baseline
base_wb = load_workbook(M.build(WORK)["path"])
ws0 = base_wb["Forecast"]
R = {lbl: row_of(ws0, lbl) for lbl in
     ["Payroll and payroll taxes", "Rent", "Owner draw", "Insurance renewal",
      "Collections on receivables", "Total money out", "Ending cash"]}
base = recalc(WORK)
print("baseline")
pay0 = [num(base, c) for c in week_cells(R["Payroll and payroll taxes"])]
check("the grid holds real numbers", any(pay0), f"payroll row {pay0}")
check("payroll is biweekly, not every week",
      0 < sum(1 for v in pay0 if v) < M.WK, f"{sum(1 for v in pay0 if v)} of {M.WK} weeks")

# Every line, against the cadence worked out independently in Python. The sheet and the
# builder have to agree on which weeks get paid, or the workbook is confidently wrong.
wrong = []
for spec in M.expand(M.CONFIG["inflows"]) + M.expand(M.CONFIG["outflows"]):
    rr = row_of(ws0, spec["label"])
    hits = M.schedule(spec)
    got = [num(base, c) for c in week_cells(rr)]
    want = [spec["amount"] if (hits == "weekly" or i in hits) else 0 for i in range(M.WK)]
    if got != want:
        wrong.append(f"{spec['label']}: want {want}, got {got}")
check("every line pays in the weeks its cadence says", not wrong, "\n        ".join(wrong))

# ---------------------------------------------------------------- 1. amount
print("\nchanging an amount")
v = build_and_read(lambda ws: {f"B{R['Payroll and payroll taxes']}": 50_000})
pay = [num(v, c) for c in week_cells(R["Payroll and payroll taxes"])]
check("a new payroll amount reaches every week it is paid",
      set(x for x in pay if x) == {50_000}, str(pay))
check("the weeks it is not paid stay at zero",
      [bool(x) for x in pay] == [bool(x) for x in pay0], str(pay))

# ---------------------------------------------------------------- 2. cadence
print("\nchanging how often")
v = build_and_read(lambda ws: {f"C{R['Rent']}": M.CAD_WEEKLY})
rent = [num(v, c) for c in week_cells(R["Rent"])]
check("switching rent to weekly fills all thirteen weeks",
      all(rent) and len(set(rent)) == 1, str(rent))

v = build_and_read(lambda ws: {f"C{R['Payroll and payroll taxes']}": M.CAD_MONTH})
pay_m = [num(v, c) for c in week_cells(R["Payroll and payroll taxes"])]
check("switching payroll to monthly thins it out",
      0 < sum(1 for x in pay_m if x) < sum(1 for x in pay0 if x),
      f"{sum(1 for x in pay_m if x)} weeks against {sum(1 for x in pay0 if x)}")

# ---------------------------------------------------------------- 3. re-timing
print("\nmoving a one-off two weeks later")
ins_row = R["Insurance renewal"]
before = [i for i, c in enumerate(week_cells(ins_row)) if num(base, c)]
v = build_and_read(lambda ws: {f"D{ins_row}": ws[f"D{ins_row}"].value + timedelta(days=14)})
after = [i for i, c in enumerate(week_cells(ins_row)) if num(v, c)]
check("the payment lands two columns later",
      len(before) == 1 and after == [before[0] + 2], f"{before} then {after}")

# ---------------------------------------------------------------- 4. ending a line
print("\nending a line early")
draw_row = R["Owner draw"]
draws0 = sum(1 for c in week_cells(draw_row) if num(base, c))
v = build_and_read(lambda ws: {f"E{draw_row}": date(2026, 9, 30)})
draws = sum(1 for c in week_cells(draw_row) if num(v, c))
check("an end date stops the line", 0 < draws < draws0, f"{draws} against {draws0}")

# ---------------------------------------------------------------- 5. the calendar
print("\nmoving the forecast start date")
v = build_and_read(lambda ws: {"B7": date(2026, 9, 7)})
h0 = num(base, f"{M.col(0)}{M.R_HEAD}")
h1 = num(v, f"{M.col(0)}{M.R_HEAD}")
check("week 1 moves with the input cell", h1 == h0 + 28, f"{h0} then {h1}")
check("week 13 follows it",
      num(v, f"{M.col(M.WK - 1)}{M.R_HEAD}") == h1 + 7 * (M.WK - 1))
pay_s = [num(v, c) for c in week_cells(R["Payroll and payroll taxes"])]
# 28 days is exactly two biweekly cycles, so an unchanged pattern is the right answer here.
check("a four-week shift leaves the biweekly pattern alone", pay_s == pay0, str(pay_s))

# This is the one that separates a real model from a table. A shift of one week flips which
# weeks payroll lands in, because the anchor is a date. A grid keyed to column position, or
# one Python filled in at build time, would not notice.
v = build_and_read(lambda ws: {"B7": date(2026, 8, 17)})
pay_o = [num(v, c) for c in week_cells(R["Payroll and payroll taxes"])]
check("a one-week shift flips which weeks payroll lands in",
      [bool(x) for x in pay_o] == [not bool(x) for x in pay0], f"{pay0} then {pay_o}")

# ---------------------------------------------------------------- 6. it all rolls up
print("\nthe rollup follows")
v = build_and_read(lambda ws: {f"B{R['Collections on receivables']}": 60_000})
end0 = num(base, f"{M.col(M.WK - 1)}{R['Ending cash']}")
end1 = num(v, f"{M.col(M.WK - 1)}{R['Ending cash']}")
check("more collections means more cash at week 13", end1 > end0, f"{end0} then {end1}")
check("the lift is the change times the weeks it applies to",
      end1 - end0 == 24_000 * M.WK, f"{end1 - end0}")

print(f"\nmodel_test: {len(fails)} fail(s)" if fails else "\nmodel_test: all green")
sys.exit(1 if fails else 0)
